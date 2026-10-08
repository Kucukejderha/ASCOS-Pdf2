"""Sembol fontu madde işaretlerinin Word çıktısındaki bozuk karşılıklarını düzeltir.

Bazı PDF'lerde madde işaretleri (madde numaralandırma noktaları) özel kodlamalı
sembol fontlarıyla çizilir; metin katmanında ise ham ASCII karakterler (ör. "G")
kalır. pdf2docx bu karakterleri tanımadığı için Word belgesine de "G" olarak
aktarılır. Bu modül, dönüşüm sonrası docx üzerinde:

1. Adı sembol/dingbat ailesine benzeyen run'lardaki bilinen karakterleri
   gerçek Unicode madde işaretlerine eşler (Wingdings/Webdings tablosu).
2. Metin deseni güçlü biçimde madde işareti olduğunu gösteren durumlarda
   (paragraf/satır başında "G" + Büyükharf+Küçükharf veya "G" + "3." gibi)
   ham karakteri "●" ile değiştirir.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.text.paragraph import Paragraph

SYMBOL_FONT_KEYWORDS = (
    "wingding",
    "webding",
    "symbol",
    "dingbat",
    "monotype sorts",
    "zapf",
    "mtextra",
    "marlett",
)

DEFAULT_BULLET = "●"

WINGDING_MAP: dict[str, str] = {
    "l": "●",
    "n": "■",
    "m": "○",
    "o": "□",
    "p": "❑",
    "q": "❒",
    "r": "❖",
    "u": "◆",
    "v": "❖",
    "Ø": "➢",
    "ü": "✓",
    "ý": "✗",
    "§": "▪",
    "·": "•",
}

_SENTENCE_END = ".!?:"

_START_CAP = re.compile(r"^G(?=[A-ZÇĞİÖŞÜ][a-zçğıöşü])")
_START_NUM = re.compile(r"^G(?=\d+\.)")
_AFTER_END = re.compile(
    r"(?:(?<=[.!?:])|(?<=[.!?:] )|(?<=\n))G(?=[A-ZÇĞİÖŞÜ][a-zçğıöşü]|\d+\.)"
)


def _is_symbol_font(name: str | None) -> bool:
    if not name:
        return False
    lowered = name.lower()
    return any(keyword in lowered for keyword in SYMBOL_FONT_KEYWORDS)


def _map_symbol_run_text(text: str) -> str:
    out: list[str] = []
    for ch in text:
        if ch in WINGDING_MAP:
            out.append(WINGDING_MAP[ch])
        elif 0xF000 <= ord(ch) < 0xF100:
            out.append(WINGDING_MAP.get(chr(ord(ch) - 0xF000), ch))
        else:
            out.append(ch)
    return "".join(out)


def _next_text(paragraph: Paragraph, run_index: int) -> str:
    runs = paragraph.runs
    for i in range(run_index + 1, len(runs)):
        if runs[i].text.strip():
            return runs[i].text.lstrip()
    return ""


def _prev_text(paragraph: Paragraph, run_index: int) -> str:
    runs = paragraph.runs
    for i in range(run_index - 1, -1, -1):
        if runs[i].text.strip():
            return runs[i].text.rstrip()
    return ""


def _follows_marker_pattern(next_text: str) -> bool:
    return bool(
        re.match(r"^[A-ZÇĞİÖŞÜ][a-zçğıöşü]", next_text)
        or re.match(r"^\d+\.", next_text)
    )


def _fix_paragraph(paragraph: Paragraph) -> int:
    count = 0

    for run in paragraph.runs:
        if _is_symbol_font(run.font.name):
            mapped = _map_symbol_run_text(run.text)
            stripped = run.text.strip()
            if mapped == run.text and 0 < len(stripped) <= 2 and stripped:
                mapped = DEFAULT_BULLET * len(stripped)
            if mapped != run.text:
                run.text = mapped
                count += 1

    runs = paragraph.runs
    first_non_empty = next((i for i, r in enumerate(runs) if r.text.strip()), None)
    for index, run in enumerate(runs):
        if run.text.strip() == "":
            continue
        text = run.text

        if text.strip() == "G" and len(runs) > 1:
            prev = _prev_text(paragraph, index)
            nxt = _next_text(paragraph, index)
            at_start = prev == ""
            after_end = bool(prev) and prev[-1] in _SENTENCE_END
            if (at_start or after_end) and _follows_marker_pattern(nxt):
                run.text = text.replace("G", DEFAULT_BULLET, 1)
                count += 1
                continue

        if index == first_non_empty:
            new_text = _START_CAP.sub(DEFAULT_BULLET, text, count=1)
            new_text = _START_NUM.sub(DEFAULT_BULLET, new_text, count=1)
        else:
            new_text = _AFTER_END.sub(DEFAULT_BULLET, text)
        if new_text != text:
            run.text = new_text
            count += 1

    return count


def fix_symbol_bullets(docx_path: str | Path) -> int:
    """Word belgesindeki sembol madde işareti bozulmalarını düzeltir.

    Değişiklik sayısını döndürür; değişiklik yoksa dosyayı yeniden yazmaz.
    """
    path = Path(docx_path)
    doc = Document(str(path))
    count = 0
    for paragraph in doc.paragraphs:
        count += _fix_paragraph(paragraph)
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    count += _fix_paragraph(paragraph)
    if count:
        doc.save(str(path))
    return count
