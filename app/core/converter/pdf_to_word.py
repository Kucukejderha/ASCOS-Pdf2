from __future__ import annotations

import os
import tempfile
import time
from pathlib import Path
from typing import Callable

from docx import Document
from pdf2docx import Converter

from app.core.models import ConversionOptions, ConversionResult, ConvertMode
from app.core.utils import unique_path

from .docx_fixups import fix_symbol_bullets
from .ocr import OcrError, get_engine, ocr_page
from .scanned_detector import analyze_pdf
from .searchable_pdf import build_searchable_pdf

ProgressCb = Callable[[float, str], None]


def _ensure_parent(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)


def _emit(progress: ProgressCb | None, fraction: float, message: str) -> None:
    if progress:
        progress(max(0.0, min(1.0, fraction)), message)


def _ocr_text_to_docx(
    src: Path,
    dst: Path,
    selected: list[int],
    options: ConversionOptions,
    progress: ProgressCb | None,
    result: ConversionResult,
) -> None:
    import pymupdf

    engine = get_engine(options.ocr_engine)
    doc = Document()
    with pymupdf.open(str(src)) as pdf:
        total = len(selected)
        for i, page_index in enumerate(selected):
            if i > 0:
                doc.add_page_break()
            page = pdf[page_index]
            text = ocr_page(page, engine, lang=options.ocr_lang)
            if text.strip():
                for line in text.splitlines():
                    doc.add_paragraph(line)
            else:
                doc.add_paragraph("")
                result.warnings.append(f"Sayfa {page_index + 1}: OCR metni boş döndü")
            result.ocr_pages.append(page_index)
            _emit(
                progress,
                0.1 + 0.85 * (i + 1) / max(total, 1),
                f"OCR: sayfa {page_index + 1}/{len(pdf)} ({engine.name})",
            )
    doc.save(str(dst))
    result.warnings.append(f"Taranmış belge: metin {engine.name} ile OCR edilerek aktarıldı")


def _append_ocr_sections(
    dst: Path,
    src: Path,
    scanned_pages: list[int],
    options: ConversionOptions,
    progress: ProgressCb | None,
    result: ConversionResult,
    fraction_start: float,
    fraction_span: float,
) -> None:
    import pymupdf

    engine = get_engine(options.ocr_engine)
    doc = Document(str(dst))
    with pymupdf.open(str(src)) as pdf:
        total = len(scanned_pages)
        for i, page_index in enumerate(scanned_pages):
            doc.add_page_break()
            doc.add_heading(f"OCR metni — Sayfa {page_index + 1}", level=1)
            text = ocr_page(pdf[page_index], engine, lang=options.ocr_lang)
            if text.strip():
                for line in text.splitlines():
                    doc.add_paragraph(line)
            else:
                result.warnings.append(f"Sayfa {page_index + 1}: OCR metni boş döndü")
            result.ocr_pages.append(page_index)
            _emit(
                progress,
                fraction_start + fraction_span * (i + 1) / max(total, 1),
                f"OCR: sayfa {page_index + 1} ({engine.name})",
            )
    doc.save(str(dst))


def _convert_via_searchable(
    src: Path,
    dst: Path,
    selected: list[int],
    scanned_pages: list[int],
    options: ConversionOptions,
    progress: ProgressCb | None,
    result: ConversionResult,
) -> bool:
    """Taranmış sayfaları aranabilir PDF'e çevirip pdf2docx ile dönüştürür.

    Başarılıysa True döner; herhangi bir adım başarısız olursa uyarı ekleyip
    False döner ve çağıran taraf eski düz metin OCR yoluna düşer.
    """
    try:
        engine = get_engine(options.ocr_engine)
    except OcrError as exc:
        result.warnings.append(f"OCR atlandı: {exc}")
        return False

    tmp = Path(tempfile.gettempdir()) / (
        f"pdf2_searchable_{os.getpid()}_{int(time.time() * 1000)}.pdf"
    )
    try:
        _emit(progress, 0.05, "Taranmış sayfalar aranabilir hale getiriliyor")
        build_searchable_pdf(
            src,
            tmp,
            engine,
            pages=scanned_pages,
            lang=options.ocr_lang,
            progress=progress,
        )
        _emit(progress, 0.92, "Word'e dönüştürülüyor")
        cv = Converter(str(tmp))
        try:
            cv.convert(str(dst), pages=selected)
        finally:
            cv.close()
        for page_index in scanned_pages:
            if page_index not in result.ocr_pages:
                result.ocr_pages.append(page_index)
        result.warnings.append(
            f"Taranmış sayfalar {engine.name} ile OCR edildi; "
            "tablolar ve düzen korunarak aktarıldı"
        )
        return True
    except Exception as exc:  # noqa: BLE001
        result.warnings.append(
            f"Düzen korumalı OCR yolu uygulanamadı ({exc}); düz metin OCR'a geçildi"
        )
        return False
    finally:
        try:
            tmp.unlink(missing_ok=True)
        except OSError:
            pass


def convert_pdf_to_word(
    src: str | Path,
    dst: str | Path,
    options: ConversionOptions | None = None,
    progress: ProgressCb | None = None,
) -> ConversionResult:
    options = options or ConversionOptions(mode=ConvertMode.WORD)
    src, dst = Path(src), Path(dst)
    if not options.overwrite:
        dst = unique_path(dst)
    _ensure_parent(dst)

    start = time.perf_counter()
    result = ConversionResult(source=src, output=dst, mode=ConvertMode.WORD)

    _emit(progress, 0.02, "Belge inceleniyor")
    analysis = analyze_pdf(src)
    selected = options.page_list(analysis.page_count) or list(range(analysis.page_count))
    if not selected:
        raise ValueError("Seçili sayfa yok")

    scanned_selected = [p for p in selected if analysis.pages[p].is_scanned]
    digital_selected = [p for p in selected if not analysis.pages[p].is_scanned]

    all_scanned = not digital_selected
    mixed = bool(scanned_selected) and bool(digital_selected)

    layout_done = False
    if scanned_selected and options.ocr and options.layout_preserve:
        _emit(progress, 0.04, "Taranmış sayfalar için düzen korumalı OCR")
        layout_done = _convert_via_searchable(
            src, dst, selected, scanned_selected, options, progress, result
        )

    if layout_done:
        pass
    elif all_scanned and options.ocr:
        _emit(progress, 0.1, "Taranmış belge algılandı, OCR başlıyor")
        _ocr_text_to_docx(src, dst, selected, options, progress, result)
    elif all_scanned and not options.ocr:
        result.warnings.append("Taranmış belge ve OCR kapalı: sayfa görüntüleri Word'e aktarılıyor")
        _emit(progress, 0.1, "Word'e dönüştürülüyor (OCR kapalı)")
        cv = Converter(str(src))
        try:
            cv.convert(str(dst), pages=selected)
        finally:
            cv.close()
    else:
        _emit(progress, 0.1, "Word'e dönüştürülüyor")
        cv = Converter(str(src))
        try:
            cv.convert(str(dst), pages=selected)
        finally:
            cv.close()
        if mixed and options.ocr:
            try:
                _append_ocr_sections(
                    dst, src, scanned_selected, options, progress, result,
                    fraction_start=0.6, fraction_span=0.35,
                )
                result.warnings.append(
                    "Karma belge: taranmış sayfalar OCR metni olarak sona eklendi"
                )
            except OcrError as exc:
                result.warnings.append(f"OCR atlandı: {exc}")

    _emit(progress, 0.98, "Kaydediliyor")
    try:
        fixed = fix_symbol_bullets(dst)
        if fixed:
            result.warnings.append(f"{fixed} bozuk madde işareti düzeltildi")
    except Exception as exc:  # noqa: BLE001
        result.warnings.append(f"Madde işareti düzeltmesi uygulanamadı: {exc}")
    result.elapsed_s = time.perf_counter() - start
    _emit(progress, 1.0, "Tamamlandı")
    return result
