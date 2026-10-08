from __future__ import annotations

import re
import time
from pathlib import Path
from typing import Callable

import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.models import ConversionOptions, ConversionResult, ConvertMode, ExcelLayout
from app.core.utils import parse_number, unique_path

from .ocr import OcrError, get_engine, ocr_page
from .scanned_detector import analyze_pdf

ProgressCb = Callable[[float, str], None]

HEADER_FILL = PatternFill("solid", fgColor="E3E6EA")
HEADER_FONT = Font(bold=True)
DIVIDER_FONT = Font(bold=True, color="555555")


def _emit(progress: ProgressCb | None, fraction: float, message: str) -> None:
    if progress:
        progress(max(0.0, min(1.0, fraction)), message)


def _sanitize_sheet_title(title: str, used: set[str]) -> str:
    clean = re.sub(r"[\[\]:*?/\\]", " ", title).strip() or "Sayfa"
    clean = clean[:31]
    candidate = clean
    i = 2
    while candidate in used:
        suffix = f" ({i})"
        candidate = clean[: 31 - len(suffix)] + suffix
        i += 1
    used.add(candidate)
    return candidate


def _table_score(tables: list[list[list[str]]]) -> int:
    return sum(
        1
        for table in tables
        for row in table
        for cell in row
        if cell is not None and str(cell).strip()
    )


def _extract_page_tables(page, pdf_path: Path, page_index: int) -> list[list[list[str]]]:
    pymupdf_tables: list[list[list[str]]] = []
    try:
        found = page.find_tables()
        for t in found.tables:
            data = t.extract()
            if data and any(any(c for c in r if c) for r in data):
                pymupdf_tables.append([[(c or "").strip() for c in r] for r in data])
    except Exception:
        pass

    pdfplumber_tables = _pdfplumber_tables(pdf_path, page_index)
    if _table_score(pdfplumber_tables) > _table_score(pymupdf_tables):
        return pdfplumber_tables
    return pymupdf_tables


def _pdfplumber_tables(pdf_path: Path, page_index: int) -> list[list[list[str]]]:
    try:
        import pdfplumber

        with pdfplumber.open(str(pdf_path)) as pdf:
            if page_index >= len(pdf.pages):
                return []
            tables = pdf.pages[page_index].extract_tables()
            out: list[list[list[str]]] = []
            for table in tables:
                rows = [[(cell or "").strip() for cell in row] for row in table]
                if any(any(cell for cell in row) for row in rows):
                    out.append(rows)
            return out
    except Exception:
        return []


def _write_table(
    ws,
    table: list[list[str]],
    start_row: int,
    with_header: bool = True,
) -> int:
    row = start_row
    col_count = max((len(r) for r in table), default=0)
    for r_index, raw_row in enumerate(table):
        for c_index in range(col_count):
            value = raw_row[c_index] if c_index < len(raw_row) else ""
            cell = ws.cell(row=row, column=c_index + 1)
            parsed = parse_number(value)
            cell.value = parsed
            if with_header and r_index == 0:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
        row += 1
    return row


def _autosize_columns(ws, widths: dict[int, int]) -> None:
    for col, width in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 9), 60)


def _collect_widths(table: list[list[str]], widths: dict[int, int]) -> None:
    for row in table:
        for c_index, cell in enumerate(row):
            length = len(str(cell)) if cell is not None else 0
            widths[c_index + 1] = max(widths.get(c_index + 1, 0), length)


def convert_pdf_to_excel(
    src: str | Path,
    dst: str | Path,
    options: ConversionOptions | None = None,
    progress: ProgressCb | None = None,
) -> ConversionResult:
    options = options or ConversionOptions(mode=ConvertMode.EXCEL)
    src, dst = Path(src), Path(dst)
    if not options.overwrite:
        dst = unique_path(dst)
    dst.parent.mkdir(parents=True, exist_ok=True)

    start = time.perf_counter()
    result = ConversionResult(source=src, output=dst, mode=ConvertMode.EXCEL)

    _emit(progress, 0.02, "Belge inceleniyor")
    analysis = analyze_pdf(src)
    selected = options.page_list(analysis.page_count) or list(range(analysis.page_count))
    if not selected:
        raise ValueError("Seçili sayfa yok")

    engine = None
    if options.ocr and any(analysis.pages[p].is_scanned for p in selected):
        try:
            engine = get_engine(options.ocr_engine)
        except OcrError as exc:
            result.warnings.append(f"OCR atlandı: {exc}")

    wb = Workbook()
    wb.remove(wb.active)
    used_titles: set[str] = set()
    all_widths: dict[int, int] = {}

    if options.excel_layout == ExcelLayout.ALL_IN_ONE:
        ws = wb.create_sheet(_sanitize_sheet_title("Tüm Tablolar", used_titles))
        row = 1
    else:
        ws = None
        row = 1

    with pymupdf.open(str(src)) as pdf:
        total = len(selected)
        for i, page_index in enumerate(selected):
            page = pdf[page_index]
            tables: list[list[list[str]]] = []
            is_scanned = analysis.pages[page_index].is_scanned

            if is_scanned:
                if engine is not None:
                    text = ocr_page(page, engine, lang=options.ocr_lang)
                    lines = [line for line in text.splitlines() if line.strip()]
                    if lines:
                        tables = [[[line] for line in lines]]
                        result.ocr_pages.append(page_index)
                        result.warnings.append(
                            f"Sayfa {page_index + 1}: taranmış sayfa OCR ile metne çevrildi"
                        )
                else:
                    result.warnings.append(
                        f"Sayfa {page_index + 1}: taranmış sayfa atlandı (OCR kapalı)"
                    )
            else:
                tables = _extract_page_tables(page, src, page_index)
                if not tables and options.text_fallback:
                    lines = [line for line in (page.get_text("text") or "").splitlines() if line.strip()]
                    if lines:
                        tables = [[[line] for line in lines]]
                        result.warnings.append(
                            f"Sayfa {page_index + 1}: tablo bulunamadı, metin tek sütun olarak aktarıldı"
                        )

            if options.excel_layout == ExcelLayout.PAGE_PER_SHEET:
                ws = wb.create_sheet(_sanitize_sheet_title(f"Sayfa {page_index + 1}", used_titles))
                row = 1

            if not tables:
                ws.cell(row=row, column=1, value=f"Sayfa {page_index + 1}: içerik bulunamadı").font = Font(italic=True)
                row += 2
            else:
                for table in tables:
                    row = _write_table(ws, table, row, with_header=True)
                    _collect_widths(table, all_widths)
                    row += 1

            _emit(
                progress,
                0.05 + 0.9 * (i + 1) / max(total, 1),
                f"Excel: sayfa {page_index + 1}/{analysis.page_count}",
            )

    if options.excel_layout == ExcelLayout.ALL_IN_ONE or wb.sheetnames:
        for sheet in wb.worksheets:
            _autosize_columns(sheet, all_widths)

    if not wb.sheetnames:
        wb.create_sheet("Boş")

    wb.save(str(dst))
    result.elapsed_s = time.perf_counter() - start
    _emit(progress, 1.0, "Tamamlandı")
    return result
