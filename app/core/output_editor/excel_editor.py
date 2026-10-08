from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.core.utils import parse_number


class ExcelWorkbookEditor:
    """Excel çalışma kitabında hücre düzeyinde düzenleme."""

    MAX_ROWS = 2000
    MAX_COLS = 120

    def __init__(self, path: str | Path) -> None:
        self.path = Path(path)
        self.wb = load_workbook(str(self.path))

    @property
    def sheet_names(self) -> list[str]:
        return list(self.wb.sheetnames)

    def displayed_shape(self, sheet: str) -> tuple[int, int]:
        ws = self.wb[sheet]
        return min(ws.max_row or 1, self.MAX_ROWS), min(ws.max_column or 1, self.MAX_COLS)

    def used_shape(self, sheet: str) -> tuple[int, int]:
        ws = self.wb[sheet]
        return ws.max_row or 1, ws.max_column or 1

    def cell_text(self, sheet: str, row: int, col: int) -> str:
        value = self.wb[sheet].cell(row=row, column=col).value
        return "" if value is None else str(value)

    def set_cell_text(self, sheet: str, row: int, col: int, text: str) -> None:
        self.wb[sheet].cell(row=row, column=col).value = parse_number(text)

    def save(self, dst: str | Path | None = None) -> Path:
        target = Path(dst) if dst else self.path
        target.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(str(target))
        return target
