from __future__ import annotations

from pathlib import Path

from docx import Document

from app.core.converter.docx_fixups import fix_symbol_bullets


def _save(doc: Document, path: Path) -> Path:
    doc.save(str(path))
    return path


def test_fixes_run_split_bullet(tmp_path: Path) -> None:
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("G")
    run.font.name = "CustomBullet"
    p.add_run("Ana Grup ve Alt Kırılımlar")
    path = _save(doc, tmp_path / "a.docx")

    count = fix_symbol_bullets(path)
    assert count == 1
    assert Document(str(path)).paragraphs[0].text == "●Ana Grup ve Alt Kırılımlar"


def test_fixes_in_run_prefix(tmp_path: Path) -> None:
    doc = Document()
    doc.add_paragraph("GSektörel Genişleme")
    doc.add_paragraph("G3. Vadeli Satış Oranları")
    path = _save(doc, tmp_path / "b.docx")

    count = fix_symbol_bullets(path)
    assert count == 2
    texts = [p.text for p in Document(str(path)).paragraphs]
    assert texts[0] == "●Sektörel Genişleme"
    assert texts[1] == "●3. Vadeli Satış Oranları"


def test_fixes_mid_paragraph_after_sentence(tmp_path: Path) -> None:
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("Değişkenlerin çokluğu nedeniyle bu kalem iptal edilir.")
    p.add_run("G")
    p.add_run("4. Kredi Kartı / Taksit Komisyonları")
    path = _save(doc, tmp_path / "c.docx")

    count = fix_symbol_bullets(path)
    assert count == 1
    text = Document(str(path)).paragraphs[0].text
    assert text == "Değişkenlerin çokluğu nedeniyle bu kalem iptal edilir.●4. Kredi Kartı / Taksit Komisyonları"


def test_does_not_touch_legitimate_text(tmp_path: Path) -> None:
    doc = Document()
    doc.add_paragraph("GPU fiyatı güncellendi.")
    doc.add_paragraph("Genel Müdür açıklama yaptı.")
    doc.add_paragraph("G3 modeli stokta.")
    p = doc.add_paragraph()
    p.add_run("Rapor").bold = True
    path = _save(doc, tmp_path / "d.docx")

    count = fix_symbol_bullets(path)
    assert count == 0
    texts = [p.text for p in Document(str(path)).paragraphs]
    assert texts == [
        "GPU fiyatı güncellendi.",
        "Genel Müdür açıklama yaptı.",
        "G3 modeli stokta.",
        "Rapor",
    ]


def test_maps_symbol_font_characters(tmp_path: Path) -> None:
    doc = Document()
    p = doc.add_paragraph()
    run = p.add_run("l")
    run.font.name = "Wingdings"
    run2 = p.add_run("n")
    run2.font.name = "ABCDEE+Wingdings"
    p.add_run("Kutu")
    path = _save(doc, tmp_path / "e.docx")

    count = fix_symbol_bullets(path)
    assert count == 2
    text = Document(str(path)).paragraphs[0].text
    assert text == "●■Kutu"


def test_fixes_first_non_empty_run(tmp_path: Path) -> None:
    doc = Document()
    p = doc.add_paragraph()
    p.add_run("")
    p.add_run("GSektörel Genişleme")
    path = _save(doc, tmp_path / "f.docx")

    count = fix_symbol_bullets(path)
    assert count == 1
    assert Document(str(path)).paragraphs[0].text == "●Sektörel Genişleme"


def test_no_change_no_rewrite(tmp_path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Sadece normal bir paragraf.")
    path = _save(doc, tmp_path / "g.docx")
    before = path.read_bytes()

    count = fix_symbol_bullets(path)
    assert count == 0
    assert path.read_bytes() == before
