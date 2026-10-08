from __future__ import annotations

from pathlib import Path

from docx import Document

from app.core.models import ConversionOptions, ConvertMode
from app.core.pipeline import convert


def _docx_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def test_digital_pdf_to_word(digital_pdf: Path, tmp_path: Path) -> None:
    result = convert(
        digital_pdf,
        ConversionOptions(mode=ConvertMode.WORD),
        output_dir=tmp_path,
    )
    assert result.output.exists()
    assert result.output.suffix == ".docx"
    assert not result.ocr_pages
    text = _docx_text(result.output)
    assert "MERHABA" in text.upper()
    assert "IKINCI" in text.upper()


def test_mixed_pdf_to_word_without_ocr(mixed_pdf: Path, tmp_path: Path) -> None:
    result = convert(
        mixed_pdf,
        ConversionOptions(mode=ConvertMode.WORD, ocr=False),
        output_dir=tmp_path,
    )
    assert result.output.exists()
    assert "DIGITAL" in _docx_text(result.output).upper()


def test_page_selection_to_word(digital_pdf: Path, tmp_path: Path) -> None:
    result = convert(
        digital_pdf,
        ConversionOptions(mode=ConvertMode.WORD, pages="2"),
        output_dir=tmp_path,
    )
    text = _docx_text(result.output).upper()
    assert "IKINCI" in text
    assert "MERHABA" not in text
