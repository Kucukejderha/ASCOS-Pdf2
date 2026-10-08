from __future__ import annotations

import os
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtWidgets import QApplication  # noqa: E402

from app.ui.main_window import MainWindow  # noqa: E402

_app: QApplication | None = None


def _get_app() -> QApplication:
    global _app
    if _app is None:
        _app = QApplication.instance() or QApplication([])
    return _app


def test_main_window_builds() -> None:
    app = _get_app()
    window = MainWindow()
    assert window.windowTitle().startswith("Pdf2")
    assert window.tabs.count() == 3
    assert window.tabs.tabText(0) == "Dönüştür"
    assert window.tabs.tabText(1) == "PDF Düzenle"
    assert window.tabs.tabText(2) == "Çıktı Düzenle"
    window.close()
    app.processEvents()


def test_convert_tab_file_management(digital_pdf: Path) -> None:
    app = _get_app()
    window = MainWindow()
    tab = window.convert_tab

    tab.add_files([str(digital_pdf)])
    tab.add_files([str(digital_pdf)])
    assert tab.file_list.count() == 1
    assert "1 dosya" in tab.file_count.text()

    tab.add_files([str(digital_pdf.with_name("ikinci.pdf"))])
    assert tab.file_list.count() == 2

    tab._clear_files()
    assert tab.file_list.count() == 0
    assert "eklenmedi" in tab.file_count.text()

    window.close()
    app.processEvents()


def test_convert_tab_mode_toggles_excel_layout() -> None:
    app = _get_app()
    window = MainWindow()
    tab = window.convert_tab

    assert tab.excel_layout_combo.isHidden()
    tab.radio_excel.setChecked(True)
    assert not tab.excel_layout_combo.isHidden()
    tab.radio_word.setChecked(True)
    assert tab.excel_layout_combo.isHidden()

    window.close()
    app.processEvents()


def test_pdf_editor_tab_opens_and_renders(digital_pdf: Path) -> None:
    app = _get_app()
    window = MainWindow()
    tab = window.pdf_editor_tab

    tab.open_path(digital_pdf)
    assert tab.session is not None
    assert tab.page_label.text() == "1 / 2"
    pixmap = tab.page_view.pixmap()
    assert pixmap is not None and not pixmap.isNull()

    tab._rotate(90)
    assert tab.session.doc[0].rotation == 90

    tab._go_to(1)
    assert tab.page_label.text() == "2 / 2"

    tab.close_session()
    window.close()
    app.processEvents()


def test_output_editor_tab_word(tmp_path: Path) -> None:
    from docx import Document

    path = tmp_path / "belge.docx"
    doc = Document()
    doc.add_paragraph("Merhaba dünya")
    doc.save(str(path))

    app = _get_app()
    window = MainWindow()
    tab = window.output_editor_tab
    tab.open_path(path)
    assert tab.word_editor is not None
    assert tab.paragraph_list.count() == 1

    tab.word_editor.set_paragraph_text(0, "Düzeltildi")
    tab._set_dirty(True)
    tab._save()
    assert not tab._dirty

    reopened = Document(str(path))
    assert reopened.paragraphs[0].text == "Düzeltildi"

    window.close()
    app.processEvents()


def test_output_editor_tab_excel(tmp_path: Path) -> None:
    from openpyxl import Workbook, load_workbook

    path = tmp_path / "kitap.xlsx"
    wb = Workbook()
    wb.active["A1"] = "Urun"
    wb.save(str(path))

    app = _get_app()
    window = MainWindow()
    tab = window.output_editor_tab
    tab.open_path(path)
    assert tab.excel_editor is not None
    assert tab.excel_grid.item(0, 0).text() == "Urun"

    tab.excel_grid.item(0, 0).setText("Yeni ürün")
    assert tab._dirty
    tab._save()

    reopened = load_workbook(str(path))
    assert reopened.active["A1"].value == "Yeni ürün"

    window.close()
    app.processEvents()


def test_pdf_editor_undo_redo_buttons(digital_pdf: Path) -> None:
    app = _get_app()
    window = MainWindow()
    tab = window.pdf_editor_tab

    tab.open_path(digital_pdf)
    assert not tab.undo_button.isEnabled()
    assert not tab.redo_button.isEnabled()

    tab._rotate(90)
    assert tab.session.doc[0].rotation == 90
    assert tab.undo_button.isEnabled()

    tab.undo()
    assert tab.session.doc[0].rotation == 0
    assert tab.redo_button.isEnabled()

    tab.redo()
    assert tab.session.doc[0].rotation == 90

    tab.close_session()
    assert not tab.undo_button.isEnabled()
    window.close()
    app.processEvents()


def test_pdf_editor_text_selection_copy(digital_pdf: Path) -> None:
    from PySide6.QtCore import QRect

    app = _get_app()
    window = MainWindow()
    tab = window.pdf_editor_tab

    tab.open_path(digital_pdf)
    zoom = tab.zoom
    rect = QRect(
        int(50 * zoom), int(55 * zoom), int(350 * zoom), int(25 * zoom)
    )
    tab._on_selection(rect)
    assert "MERHABA" in tab._selection_text.upper()

    tab.copy_selection()
    assert "MERHABA" in QApplication.clipboard().text().upper()

    tab._clear_selection()
    assert tab._selection_text == ""

    tab.close_session()
    window.close()
    app.processEvents()


def test_pdf_editor_erase_and_redact_tools(digital_pdf: Path) -> None:
    from PySide6.QtCore import QRect

    app = _get_app()
    window = MainWindow()
    tab = window.pdf_editor_tab

    tab.open_path(digital_pdf)
    zoom = tab.zoom
    rect = QRect(int(60 * zoom), int(55 * zoom), int(240 * zoom), int(25 * zoom))

    before = len(tab.session.doc[0].get_drawings())
    tab._on_erase(rect)
    assert len(tab.session.doc[0].get_drawings()) > before

    tab._on_redact(rect)
    assert "MERHABA" not in tab.session.doc[0].get_text().upper()
    assert tab.undo_button.isEnabled()

    tab.undo()
    assert "MERHABA" in tab.session.doc[0].get_text().upper()

    tab.close_session()
    window.close()
    app.processEvents()


def test_pdf_editor_edit_text_tool(digital_pdf: Path, monkeypatch) -> None:  # noqa: ANN001
    from PySide6.QtCore import QRect
    from PySide6.QtWidgets import QInputDialog

    monkeypatch.setattr(
        QInputDialog, "getMultiLineText", lambda *a, **k: ("YENI METIN", True)
    )

    app = _get_app()
    window = MainWindow()
    tab = window.pdf_editor_tab
    tab.open_path(digital_pdf)

    zoom = tab.zoom
    rect = QRect(int(60 * zoom), int(55 * zoom), int(340 * zoom), int(25 * zoom))
    tab._on_edit(rect)

    text = tab.session.doc[0].get_text().replace("\xa0", " ").upper()
    assert "YENI METIN" in text
    assert "MERHABA" not in text
    assert tab.undo_button.isEnabled()

    tab.close_session()
    window.close()
    app.processEvents()


def test_pdf_editor_draw_tools(digital_pdf: Path) -> None:
    from PySide6.QtCore import QPoint, QRect

    app = _get_app()
    window = MainWindow()
    tab = window.pdf_editor_tab
    tab.open_path(digital_pdf)

    before = len(tab.session.doc[0].get_drawings())
    tab._on_shape_rect("rect", QRect(360, 120, 100, 70))
    tab._on_shape_line("arrow", QPoint(360, 220), QPoint(460, 260))
    tab._on_pen([QPoint(360, 300), QPoint(380, 315), QPoint(400, 305)])
    after = len(tab.session.doc[0].get_drawings())
    assert after >= before + 3

    tab.undo()
    assert len(tab.session.doc[0].get_drawings()) == after - 1

    tab.close_session()
    window.close()
    app.processEvents()


def test_pdf_editor_insert_image_tool(digital_pdf: Path, sample_png: Path, monkeypatch) -> None:  # noqa: ANN001
    from PySide6.QtCore import QRect

    app = _get_app()
    window = MainWindow()
    tab = window.pdf_editor_tab
    tab.open_path(digital_pdf)

    monkeypatch.setattr(tab, "_resolve_image_source", lambda: {"path": str(sample_png)})
    zoom = tab.zoom
    rect = QRect(int(300 * zoom), int(400 * zoom), int(200 * zoom), int(100 * zoom))

    before = len(tab.session.doc[0].get_images())
    tab._on_image_rect("image", rect)
    assert len(tab.session.doc[0].get_images()) == before + 1
    assert tab.undo_button.isEnabled()

    tab.undo()
    assert len(tab.session.doc[0].get_images()) == before

    tab.close_session()
    window.close()
    app.processEvents()


def test_pdf_editor_page_menu_has_stamp_actions(digital_pdf: Path) -> None:
    app = _get_app()
    window = MainWindow()
    tab = window.pdf_editor_tab
    tab.open_path(digital_pdf)

    labels = [action.text() for action in tab.page_menu.actions()]
    assert "Filigran ekle…" in labels
    assert "Sayfa numarası ekle…" in labels

    tab.close_session()
    window.close()
    app.processEvents()


def test_pdf_editor_thumbnails_panel(digital_pdf: Path) -> None:
    app = _get_app()
    window = MainWindow()
    tab = window.pdf_editor_tab
    tab.open_path(digital_pdf)
    assert tab.thumbnails.count() == 2

    tab._on_thumbnails_reordered([1, 0])
    assert "IKINCI" in tab.session.doc[0].get_text().upper()
    assert tab.thumbnails.count() == 2

    tab.thumbnails.setCurrentRow(0)
    tab._duplicate_selected()
    assert tab.session.page_count == 3
    assert tab.thumbnails.count() == 3

    tab._copy_selected()
    tab._paste_pages()
    assert tab.session.page_count == 4

    tab.undo()
    assert tab.session.page_count == 3

    tab.close_session()
    assert tab.thumbnails.count() == 0
    window.close()
    app.processEvents()


def test_pdf_editor_crop_tool(digital_pdf: Path, monkeypatch) -> None:  # noqa: ANN001
    from PySide6.QtCore import QRect
    from PySide6.QtWidgets import QInputDialog

    monkeypatch.setattr(
        QInputDialog, "getItem", lambda *a, **k: ("Geçerli sayfa", True)
    )

    app = _get_app()
    window = MainWindow()
    tab = window.pdf_editor_tab
    tab.open_path(digital_pdf)

    zoom = tab.zoom
    before_width = tab.session.doc[0].rect.width
    rect = QRect(int(60 * zoom), int(60 * zoom), int(300 * zoom), int(280 * zoom))

    tab._on_crop(rect)
    assert tab.session.doc[0].rect.width < before_width
    assert tab.thumbnails.count() == 2

    tab._reset_crop()
    assert abs(tab.session.doc[0].rect.width - before_width) < 1

    tab.close_session()
    window.close()
    app.processEvents()


def test_pdf_editor_security_and_form_dialogs(form_pdf: Path) -> None:
    from app.ui.security_dialogs import FormFillDialog, PasswordDialog

    app = _get_app()
    window = MainWindow()
    tab = window.pdf_editor_tab
    tab.open_path(form_pdf)

    dialog = FormFillDialog(tab.session.list_form_fields())
    assert "ad" in dialog._editors
    assert "onay" in dialog._editors

    password_dialog = PasswordDialog()
    password_dialog.user_edit.setText("deneme")
    password_dialog.print_check.setChecked(False)
    values = password_dialog.values()
    assert values["user_password"] == "deneme"
    assert values["allow_print"] is False
    assert values["allow_copy"] is True

    labels = [action.text() for action in tab.page_menu.actions()]
    assert "Form alanlarını doldur…" in labels
    assert "Şifre ekle…" in labels
    assert "Şifre korumasını kaldır" in labels

    tab.close_session()
    window.close()
    app.processEvents()


def test_conversion_finishes_and_app_stays_alive(digital_pdf: Path, tmp_path: Path) -> None:
    import time

    app = _get_app()
    window = MainWindow()
    tab = window.convert_tab
    tab.add_files([str(digital_pdf)])
    tab.output_edit.setText(str(tmp_path))

    tab._start_conversion()
    assert tab.is_busy()

    deadline = time.time() + 60
    while time.time() < deadline and (tab.is_busy() or tab._thread is not None):
        app.processEvents()
        time.sleep(0.02)

    assert tab._thread is None
    assert not tab.is_busy()
    assert tab.convert_button.isEnabled()
    assert (tmp_path / "digital.docx").exists()

    window.close()
    app.processEvents()
    assert window.tabs.count() == 3
