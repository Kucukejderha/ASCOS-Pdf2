from __future__ import annotations

from pathlib import Path

import pymupdf
import pytest
from docx import Document

from app.core.converter.ocr import get_engine
from app.core.converter.searchable_pdf import build_searchable_pdf, extract_ocr_table
from app.core.models import ConversionOptions, ConvertMode
from app.core.pipeline import convert


def _norm(text: str) -> str:
    return text.replace("\xa0", " ").upper()


def _docx_all_text(path: Path) -> str:
    doc = Document(str(path))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


@pytest.mark.slow
def test_build_searchable_pdf_has_text_and_lines(
    scanned_table_pdf: Path, tmp_path: Path
) -> None:
    mid = tmp_path / "mid.pdf"
    engine = get_engine("rapidocr")
    info = build_searchable_pdf(scanned_table_pdf, mid, engine, pages=[0])
    assert info["rebuilt_pages"] == [0]

    with pymupdf.open(str(mid)) as doc:
        text = _norm(doc[0].get_text())
        assert "URUN" in text or "KALEM" in text
        assert len(doc[0].get_drawings()) >= 4


@pytest.mark.slow
def test_scanned_table_layout_preserved_in_docx(
    scanned_table_pdf: Path, tmp_path: Path
) -> None:
    result = convert(
        scanned_table_pdf,
        ConversionOptions(mode=ConvertMode.WORD, layout_preserve=True),
        output_dir=tmp_path,
    )
    assert result.ocr_pages == [0]
    assert result.output.exists()

    doc = Document(str(result.output))
    assert len(doc.tables) >= 1
    all_text = _norm(_docx_all_text(result.output))
    assert "KALEM" in all_text
    assert "15,50" in all_text or "15.50" in all_text
    assert len(doc.tables[0].rows) >= 3
    assert len(doc.tables[0].columns) >= 3
    # Görsel referans istenmiyor: çıktı tamamen düzenlenebilir metin/tablo olmalı.
    import zipfile

    names = zipfile.ZipFile(result.output).namelist()
    assert not any("word/media/" in name for name in names)


@pytest.mark.slow
def test_scanned_table_ocr_extracts_cells(
    scanned_table_pdf: Path,
) -> None:
    import pymupdf
    from app.core.converter.ocr import render_page

    engine = get_engine("rapidocr")
    with pymupdf.open(str(scanned_table_pdf)) as pdf:
        data = extract_ocr_table(render_page(pdf[0]), engine)
    assert len(data) >= 3
    assert len(data[0]) >= 3
    cells = " ".join(cell for row in data for cell in row).upper()
    assert "KALEM" in cells


@pytest.mark.slow
def test_scanned_table_plain_ocr_fallback(
    scanned_table_pdf: Path, tmp_path: Path
) -> None:
    result = convert(
        scanned_table_pdf,
        ConversionOptions(mode=ConvertMode.WORD, layout_preserve=False),
        output_dir=tmp_path,
    )
    assert result.ocr_pages == [0]
    doc = Document(str(result.output))
    assert len(doc.tables) == 0
    assert "KALEM" in _norm(_docx_all_text(result.output))
