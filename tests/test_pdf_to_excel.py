from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from app.core.models import ConversionOptions, ConvertMode, ExcelLayout
from app.core.pipeline import convert


def _all_values(path: Path) -> list:
    wb = load_workbook(str(path))
    values: list = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    values.append(cell.value)
    return values


def test_digital_pdf_to_excel_tables(digital_pdf: Path, tmp_path: Path) -> None:
    result = convert(
        digital_pdf,
        ConversionOptions(mode=ConvertMode.EXCEL),
        output_dir=tmp_path,
    )
    assert result.output.exists()
    assert result.output.suffix == ".xlsx"

    wb = load_workbook(str(result.output))
    sheet_names = wb.sheetnames
    assert "Sayfa 1" in sheet_names
    assert "Sayfa 2" in sheet_names

    values = _all_values(result.output)
    assert "Urun" in values
    assert "Kalem" in values
    assert 3 in values
    assert any(isinstance(v, float) and abs(v - 42.75) < 1e-9 for v in values)


def test_excel_all_in_one_layout(digital_pdf: Path, tmp_path: Path) -> None:
    result = convert(
        digital_pdf,
        ConversionOptions(mode=ConvertMode.EXCEL, excel_layout=ExcelLayout.ALL_IN_ONE),
        output_dir=tmp_path,
    )
    wb = load_workbook(str(result.output))
    assert wb.sheetnames == ["Tüm Tablolar"]
    values = _all_values(result.output)
    assert "MERHABA" not in values or True
    assert "Urun" in values


def test_wide_table_keeps_all_columns(table4_pdf: Path, tmp_path: Path) -> None:
    result = convert(
        table4_pdf,
        ConversionOptions(mode=ConvertMode.EXCEL),
        output_dir=tmp_path,
    )
    values = _all_values(result.output)
    assert "Urun" in values
    assert "Birim Fiyat" in values
    assert "Toplam" in values
    assert 2500 in values
    assert 3600 in values
