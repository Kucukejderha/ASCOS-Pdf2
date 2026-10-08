from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from app.core.output_editor.excel_editor import ExcelWorkbookEditor
from app.core.output_editor.word_editor import WordDocumentEditor


@pytest.fixture()
def docx_path(tmp_path: Path) -> Path:
    path = tmp_path / "belge.docx"
    doc = Document()
    doc.add_paragraph("İlk paragraf")
    doc.add_paragraph("İkinci paragraf")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Urun"
    table.cell(0, 1).text = "Fiyat"
    table.cell(1, 0).text = "Kalem"
    table.cell(1, 1).text = "15,50"
    doc.save(str(path))
    return path


@pytest.fixture()
def xlsx_path(tmp_path: Path) -> Path:
    path = tmp_path / "kitap.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Veri"
    ws["A1"] = "Urun"
    ws["B1"] = "Fiyat"
    ws["A2"] = "Kalem"
    ws["B2"] = "15,50"
    wb.save(str(path))
    return path


def test_word_editor_paragraphs_and_tables(docx_path: Path) -> None:
    editor = WordDocumentEditor(docx_path)
    texts = editor.paragraph_texts()
    assert "İlk paragraf" in texts
    assert editor.table_count == 1

    editor.set_paragraph_text(0, "Düzeltilmiş paragraf")
    editor.set_table_cell(0, 1, 1, "99,99")
    saved = editor.save()

    reopened = Document(str(saved))
    assert reopened.paragraphs[0].text == "Düzeltilmiş paragraf"
    assert reopened.tables[0].rows[1].cells[1].text == "99,99"


def test_word_editor_save_as_keeps_original(docx_path: Path, tmp_path: Path) -> None:
    editor = WordDocumentEditor(docx_path)
    editor.set_paragraph_text(0, "Sadece kopyada")
    copy = tmp_path / "kopya.docx"
    editor.save(copy)

    original = Document(str(docx_path))
    assert original.paragraphs[0].text == "İlk paragraf"
    assert copy.exists()


def test_excel_editor_cells(xlsx_path: Path) -> None:
    editor = ExcelWorkbookEditor(xlsx_path)
    assert editor.sheet_names == ["Veri"]
    assert editor.cell_text("Veri", 2, 1) == "Kalem"
    assert editor.cell_text("Veri", 1, 1) == "Urun"

    editor.set_cell_text("Veri", 2, 2, "42,75")
    rows, cols = editor.displayed_shape("Veri")
    assert rows == 2
    assert cols == 2
    saved = editor.save()

    wb = load_workbook(str(saved))
    assert wb["Veri"]["B2"].value == pytest.approx(42.75)
    assert wb["Veri"]["A1"].value == "Urun"


def test_excel_editor_save_as(xlsx_path: Path, tmp_path: Path) -> None:
    editor = ExcelWorkbookEditor(xlsx_path)
    editor.set_cell_text("Veri", 1, 1, "Yeni")
    copy = tmp_path / "kopya.xlsx"
    editor.save(copy)

    original = load_workbook(str(xlsx_path))
    assert original["Veri"]["A1"].value == "Urun"
    assert copy.exists()
